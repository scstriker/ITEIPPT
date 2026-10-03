"""Generate synthetic E2E fixture files (all marked 示例).

Run from src/: python e2e/fixtures/make_fixtures.py
Needs python-docx, reportlab and Pillow, and the SimHei font (Windows: C:/Windows/Fonts/simhei.ttf).
"""

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).parent
TITLE = "医用呼吸机压力传感器技术要求（示例）"
PARAGRAPHS = [
    "第一章 项目背景与必要性",
    "当前国产医用呼吸机压力传感器依赖进口，供应链风险较高，需要制定统一的技术要求。",
    "本标准规定了压力传感器的测量范围为 0～10 kPa，准确度优于 1%，响应时间不超过 20 ms。",
    "第二章 主要技术内容",
    "标准制定过程包括立项、起草、征求意见和送审四个阶段，历时 18 个月。",
    "试验方法采用标准压力源比对，环境温度 23 ℃，相对湿度 50%。",
]


def make_docx() -> None:
    doc = Document()
    doc.add_paragraph(TITLE)
    for text in PARAGRAPHS:
        doc.add_paragraph(text)
    doc.save(HERE / "技术要求（示例）.docx")


def make_pdf() -> None:
    pdfmetrics.registerFont(TTFont("SimHei", "C:/Windows/Fonts/simhei.ttf"))
    c = canvas.Canvas(str(HERE / "技术要求（示例）.pdf"), pagesize=A4)
    c.setFont("SimHei", 12)
    y = 800
    for line in [TITLE, *PARAGRAPHS]:
        c.drawString(60, y, line)
        y -= 24
    c.save()

    blank = canvas.Canvas(str(HERE / "空白扫描件（示例）.pdf"), pagesize=A4)
    blank.rect(60, 600, 200, 100)
    blank.save()


def make_images() -> None:
    colors = ["#044470", "#0080cc", "#b81c27", "#8fa1b3"]
    for index, color in enumerate(colors, start=1):
        image = Image.new("RGB", (320, 180), color)
        ImageDraw.Draw(image).rectangle((20, 20, 300, 160), outline="white", width=4)
        image.save(HERE / f"配图{index}（示例）.png")


def make_text() -> None:
    (HERE / "说明（示例）.txt").write_text("这是一个文本文件（示例）。", encoding="utf-8")


if __name__ == "__main__":
    make_docx()
    make_pdf()
    make_images()
    make_text()
    print("fixtures written to", HERE)
